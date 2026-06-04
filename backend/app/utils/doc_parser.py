"""文档解析工具 - 解析 docx/xlsx 文件"""

from pathlib import Path
from docx import Document as DocxDocument
from openpyxl import load_workbook


def parse_docx(file_path: str | Path) -> list[dict]:
    """解析 docx 文件，返回段落列表

    Returns:
        [{"text": "...", "type": "paragraph/heading/table", "metadata": {...}}]
    """
    doc = DocxDocument(str(file_path))
    sections = []
    current_heading = ""

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        if para.style.name.startswith("Heading"):
            current_heading = text
            sections.append({
                "text": text,
                "type": "heading",
                "metadata": {"heading": text, "level": para.style.name},
            })
        else:
            sections.append({
                "text": text,
                "type": "paragraph",
                "metadata": {"heading": current_heading},
            })

    # 解析表格
    for table in doc.tables:
        headers = [cell.text.strip() for cell in table.rows[0].cells] if table.rows else []
        for row in table.rows[1:]:
            row_text = " | ".join(f"{h}: {cell.text.strip()}" for h, cell in zip(headers, row.cells) if cell.text.strip())
            if row_text:
                sections.append({
                    "text": row_text,
                    "type": "table",
                    "metadata": {"headers": headers},
                })

    return sections


def parse_xlsx(file_path: str | Path) -> list[dict]:
    """解析 xlsx 文件，返回行数据列表"""
    wb = load_workbook(str(file_path), read_only=True, data_only=True)
    sections = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        # 第一行作为表头
        headers = [str(h) if h else f"列{i}" for i, h in enumerate(rows[0], 1)]

        for row in rows[1:]:
            row_data = {h: str(v) if v is not None else "" for h, v in zip(headers, row)}
            row_text = " | ".join(f"{k}: {v}" for k, v in row_data.items() if v)
            if row_text:
                sections.append({
                    "text": row_text,
                    "type": "row",
                    "metadata": {"sheet": sheet_name, "headers": headers},
                })

    wb.close()
    return sections


def parse_file(file_path: str | Path) -> list[dict]:
    """根据文件类型自动选择解析器"""
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".docx":
        return parse_docx(path)
    elif suffix in (".xlsx", ".xls"):
        return parse_xlsx(path)
    elif suffix == ".txt":
        with open(path, "r", encoding="utf-8") as f:
            content = f.read()
        return [{"text": para.strip(), "type": "paragraph", "metadata": {}}
                for para in content.split("\n\n") if para.strip()]
    else:
        raise ValueError(f"不支持的文件类型: {suffix}")
